"""Tests for CvDOCX DOCX generation."""
import io

from docx import Document as DocxDocument

from app import CvDOCX


class TestCvDOCX:
    def test_get_bytes_is_bytes(self, cv):
        assert isinstance(CvDOCX(cv).get_bytes(), bytes)

    def test_get_bytes_non_empty(self, cv):
        assert len(CvDOCX(cv).get_bytes()) > 0

    def test_pk_magic_bytes(self, cv):
        assert CvDOCX(cv).get_bytes()[:2] == b"PK"

    def test_parseable_by_python_docx(self, cv):
        parsed = DocxDocument(io.BytesIO(CvDOCX(cv).get_bytes()))
        assert parsed is not None

    def test_name_in_paragraphs(self, cv):
        parsed = DocxDocument(io.BytesIO(CvDOCX(cv).get_bytes()))
        text = "\n".join(p.text for p in parsed.paragraphs)
        assert "Ada Lovelace" in text

    def test_company_in_paragraphs(self, cv):
        parsed = DocxDocument(io.BytesIO(CvDOCX(cv).get_bytes()))
        text = "\n".join(p.text for p in parsed.paragraphs)
        assert "Babbage & Co" in text

    def test_section_headings_uppercase(self, cv):
        parsed = DocxDocument(io.BytesIO(CvDOCX(cv).get_bytes()))
        text = "\n".join(p.text for p in parsed.paragraphs)
        assert "PROFILE" in text
        assert "EXPERIENCE" in text
        assert "TECHNICAL SKILLS" in text
        assert "EDUCATION" in text

    def test_sparse_cv_does_not_crash(self):
        doc = CvDOCX({"personal": {"name": "Min Person"}})
        assert doc.get_bytes()[:2] == b"PK"
