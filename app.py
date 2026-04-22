import io
import json
import os
from pathlib import Path

import anthropic
import yaml
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips
from dotenv import load_dotenv
from fpdf import FPDF
from flask import (
    Flask, Response, render_template,
    stream_with_context, request
)

load_dotenv()
app = Flask(__name__)

CONFIG_FILE = Path("cv_config.yaml")

SYSTEM_PROMPT = """You are an expert CV writer and career coach with 15+ years of experience in technical recruitment.
Your task is to tailor the candidate's existing CV for a specific role and company.

Guidelines:
- Rewrite the Profile/Summary section to directly address the role's key requirements
- Reorder and emphasise experiences and bullet points most relevant to the target role
- Mirror keywords and phrases from the job description naturally throughout
- Use strong action verbs; preserve and highlight all quantified achievements
- Do NOT fabricate or invent any experience, company, technology, or achievement
- Reorder the Skills section so the most relevant skills appear first
- Keep the same general length and structure as the original CV
- Output clean Markdown only — no preamble, no commentary, no code fences"""


# ── Config loading ────────────────────────────────────────────────────

def load_cv() -> dict:
    """Load and return the CV config from cv_config.yaml."""
    return yaml.safe_load(CONFIG_FILE.read_text(encoding="utf-8"))


def cv_to_markdown(cv: dict) -> str:
    """Convert CV config dict to Markdown text (used for AI and txt download)."""
    p = cv.get("personal", {})
    lines = [f"# {p.get('name', '')}"]
    lines.append("")

    contact = []
    if p.get("location"):  contact.append(f"**Location:** {p['location']}")
    if p.get("phone"):     contact.append(f"**Phone:** {p['phone']}")
    if p.get("email"):     contact.append(f"**Email:** {p['email']}")
    if p.get("github"):    contact.append(f"**GitHub:** github.com/{p['github']}")
    if p.get("linkedin"):  contact.append(f"**LinkedIn:** linkedin.com/in/{p['linkedin']}")
    for c in contact:
        lines.append(f"- {c}")
    lines.append("")

    if cv.get("profile"):
        lines += ["## Profile", "", cv["profile"].strip(), ""]

    if cv.get("experience"):
        lines += ["## Professional Experience", ""]
        for job in cv["experience"]:
            lines.append(f"### {job.get('company', '')}")
            lines.append(f"- **Position:** {job.get('title', '')}")
            lines.append(f"- **Period:** {job.get('period', '')}")
            if job.get("location"):
                lines.append(f"- **Location:** {job['location']}")
            lines.append("")
            for h in job.get("highlights", []):
                if h.get("title"):
                    lines.append(f"**{h['title']}**")
                for pt in h.get("points", []):
                    lines.append(f"- {' '.join(pt.split())}")
                lines.append("")
            for pt in job.get("points", []):
                lines.append(f"- {' '.join(pt.split())}")
            if job.get("points"):
                lines.append("")

    if cv.get("skills"):
        lines += ["## Technical Skills", ""]
        for s in cv["skills"]:
            lines.append(f"- **{s['category']}:** {s['items']}")
        lines.append("")

    if cv.get("education"):
        lines += ["## Education", ""]
        for e in cv["education"]:
            suffix = f" ({e['year']})" if e.get("year") else ""
            lines.append(f"- **{e['degree']}** — {e.get('institution', '')}{suffix}")

    return "\n".join(lines)


# ── PDF generation ────────────────────────────────────────────────────

def _sanitize(text: str) -> str:
    """Replace characters outside latin-1 with safe ASCII equivalents."""
    subs = {
        '\u2014': '-',   # em dash
        '\u2013': '-',   # en dash
        '\u2018': "'",   # left single quote
        '\u2019': "'",   # right single quote
        '\u201c': '"',   # left double quote
        '\u201d': '"',   # right double quote
        '\u2022': '-',   # bullet
        '\u2026': '...',  # ellipsis
        '\u00b7': '.',   # middle dot  (used in header join)
    }
    for char, rep in subs.items():
        text = text.replace(char, rep)
    return text.encode('latin-1', errors='replace').decode('latin-1')


class CvPDF(FPDF):
    """Professional CV PDF built from the cv_config dict."""

    MARGIN_LR = 18
    MARGIN_TOP = 16

    # Colours
    C_TEXT      = (26, 29, 35)
    C_MUTED     = (90, 98, 112)
    C_ACCENT    = (37, 99, 235)
    C_RULE      = (210, 215, 220)

    def __init__(self, cv: dict):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.cv = cv
        self.set_margins(self.MARGIN_LR, self.MARGIN_TOP, self.MARGIN_LR)
        self.set_auto_page_break(auto=True, margin=16)
        self.add_page()
        self._build()

    # ── helpers ──────────────────────────────────────────────────────

    def _w(self):
        return self.w - 2 * self.MARGIN_LR

    def _rule(self):
        r, g, b = self.C_RULE
        self.set_draw_color(r, g, b)
        self.set_line_width(0.3)
        x = self.MARGIN_LR
        self.line(x, self.get_y(), x + self._w(), self.get_y())
        self.ln(1.5)

    def _section_title(self, text: str):
        self.ln(4)
        r, g, b = self.C_MUTED
        self.set_text_color(r, g, b)
        self.set_font("Helvetica", "B", 7.5)
        self.cell(self._w(), 4, text.upper(), ln=True)
        self._rule()
        r, g, b = self.C_TEXT
        self.set_text_color(r, g, b)

    def _bullet(self, text: str, indent: float = 4):
        clean = _sanitize(" ".join(text.split()))
        self.set_font("Helvetica", "", 8.5)
        r, g, b = self.C_TEXT
        self.set_text_color(r, g, b)
        bw = self._w() - indent
        self.set_x(self.MARGIN_LR + indent - 3)
        self.cell(3, 4.5, "-", ln=False)
        self.set_x(self.MARGIN_LR + indent)
        self.multi_cell(bw, 4.5, clean)

    def _wrapped(self, text: str, size: float = 8.5, style: str = "", color=None):
        if color:
            self.set_text_color(*color)
        else:
            self.set_text_color(*self.C_TEXT)
        self.set_font("Helvetica", style, size)
        clean = _sanitize(" ".join(text.split()))
        self.multi_cell(self._w(), 4.8, clean)

    # ── sections ─────────────────────────────────────────────────────

    def _build(self):
        self._header_block()
        if self.cv.get("profile"):
            self._profile()
        if self.cv.get("experience"):
            self._experience()
        if self.cv.get("skills"):
            self._skills()
        if self.cv.get("education"):
            self._education()

    def _header_block(self):
        p = self.cv.get("personal", {})
        self.set_font("Helvetica", "B", 20)
        self.set_text_color(*self.C_TEXT)
        self.cell(self._w(), 9, _sanitize(p.get("name", "")), ln=True)
        if p.get("title"):
            self.set_font("Helvetica", "", 10)
            self.set_text_color(*self.C_MUTED)
            self.cell(self._w(), 5, _sanitize(p["title"]), ln=True)
        self.ln(2)
        parts = []
        if p.get("location"):  parts.append(p["location"])
        if p.get("phone"):     parts.append(p["phone"])
        if p.get("email"):     parts.append(p["email"])
        if p.get("github"):    parts.append(f"github.com/{p['github']}")
        if p.get("linkedin"):  parts.append(f"linkedin.com/in/{p['linkedin']}")
        if parts:
            self.set_font("Helvetica", "", 8)
            self.set_text_color(*self.C_MUTED)
            self.cell(self._w(), 4, _sanitize("  .  ".join(parts)), ln=True)
        self.ln(1)
        self._rule()

    def _profile(self):
        self._section_title("Profile")
        self._wrapped(self.cv["profile"].strip(), size=8.8)
        self.ln(1)

    def _experience(self):
        self._section_title("Professional Experience")
        for job in self.cv["experience"]:
            # Job title + period on same line
            y_before = self.get_y()
            self.set_font("Helvetica", "B", 9)
            self.set_text_color(*self.C_TEXT)
            period = _sanitize(job.get("period", ""))
            pw = self.get_string_width(period) + 1
            self.cell(self._w() - pw, 5, _sanitize(job.get("title", "")), ln=False)
            self.set_font("Helvetica", "", 8)
            self.set_text_color(*self.C_MUTED)
            self.cell(pw, 5, period, align="R", ln=True)

            company = job.get("company", "")
            if job.get("location"):
                company += f"  -  {job['location']}"
            self.set_font("Helvetica", "I", 8.2)
            self.set_text_color(*self.C_MUTED)
            self.cell(self._w(), 4, _sanitize(company), ln=True)
            self.ln(1)

            for h in job.get("highlights", []):
                if h.get("title"):
                    self.set_font("Helvetica", "B", 8.2)
                    self.set_text_color(*self.C_TEXT)
                    self.set_x(self.MARGIN_LR + 2)
                    self.cell(self._w() - 2, 4, _sanitize(h["title"]), ln=True)
                for pt in h.get("points", []):
                    self._bullet(pt, indent=6)
                self.ln(0.5)

            # Flat points (no sub-sections)
            for pt in job.get("points", []):
                self._bullet(pt, indent=4)

            self.ln(2)

    def _skills(self):
        self._section_title("Technical Skills")
        for s in self.cv.get("skills", []):
            cat   = s.get("category", "")
            items = str(s.get("items", ""))
            cat_w = 38
            self.set_font("Helvetica", "B", 8.2)
            self.set_text_color(*self.C_TEXT)
            x_start = self.get_x()
            y_start = self.get_y()
            self.set_x(self.MARGIN_LR)
            self.cell(cat_w, 4.5, cat + ":", ln=False)

            self.set_font("Helvetica", "", 8.2)
            self.set_text_color(*self.C_MUTED)
            self.set_x(self.MARGIN_LR + cat_w)
            self.multi_cell(self._w() - cat_w, 4.5, items)
        self.ln(1)

    def _education(self):
        self._section_title("Education")
        for e in self.cv.get("education", []):
            self.set_font("Helvetica", "B", 8.5)
            self.set_text_color(*self.C_TEXT)
            self.cell(self._w(), 4.5, _sanitize(e.get("degree", "")), ln=True)
            inst = e.get("institution", "")
            if e.get("year"):
                inst += f"  ({e['year']})"
            self.set_font("Helvetica", "", 8.2)
            self.set_text_color(*self.C_MUTED)
            self.cell(self._w(), 4, _sanitize(inst), ln=True)


# ── DOCX generation ──────────────────────────────────────────────────

class CvDOCX:
    """
    Professional CV DOCX matching the Calibri / Navy style.

    Layout mirrors the JS docx reference implementation:
      - Calibri throughout, navy (#1F3A5F) accent colour
      - Section headings: uppercase, bold 13pt, navy bottom border
      - Job title line: bold left, italic period right-aligned via tab stop
      - Company line: bold navy left, italic grey location right
      - Bullet points: hanging-indent (left=360tw, hanging=260tw)
      - Project sub-headers: bold italic grey
    """

    NAVY = RGBColor(0x1F, 0x3A, 0x5F)
    GRAY = RGBColor(0x59, 0x59, 0x59)

    def __init__(self, cv: dict):
        self.cv = cv
        self.doc = DocxDocument()
        self._setup_page()
        self._build()

    # ── page setup ────────────────────────────────────────────────────

    def _setup_page(self):
        sec = self.doc.sections[0]
        sec.page_width  = Twips(11906)   # A4 width
        sec.page_height = Twips(16838)   # A4 height
        sec.left_margin   = Twips(1000)
        sec.right_margin  = Twips(1000)
        sec.top_margin    = Twips(600)
        sec.bottom_margin = Twips(600)

    # ── low-level helpers ─────────────────────────────────────────────

    def _sp(self, para, *, before=0, after=80, line=276):
        """Set paragraph spacing (twips)."""
        fmt = para.paragraph_format
        fmt.space_before  = Twips(before)
        fmt.space_after   = Twips(after)
        fmt.line_spacing  = Twips(line)

    def _run(self, para, text, *, size=11, bold=False, italic=False, color=None):
        r = para.add_run(text)
        r.font.name  = "Calibri"
        r.font.size  = Pt(size)
        r.bold       = bold
        r.italic     = italic
        if color is not None:
            r.font.color.rgb = (
                RGBColor.from_string(color) if isinstance(color, str) else color
            )
        return r

    def _right_tab(self, para, pos: int = 9360):
        """Add a right-aligned tab stop at pos (twips)."""
        pPr  = para._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab  = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:pos"), str(pos))
        tabs.append(tab)
        pPr.append(tabs)

    def _bottom_border(self, para):
        """Add navy bottom border to paragraph (section heading style)."""
        pPr  = para._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        btm  = OxmlElement("w:bottom")
        btm.set(qn("w:val"),   "single")
        btm.set(qn("w:sz"),    "8")
        btm.set(qn("w:space"), "2")
        btm.set(qn("w:color"), "1F3A5F")
        pBdr.append(btm)
        pPr.append(pBdr)

    # ── paragraph builders ────────────────────────────────────────────

    def _para(self, text="", *, size=11, bold=False, italic=False,
              color=None, before=0, after=80, line=276, align=None):
        p = self.doc.add_paragraph()
        self._sp(p, before=before, after=after, line=line)
        if align:
            p.alignment = align
        if text:
            self._run(p, text, size=size, bold=bold, italic=italic, color=color)
        return p

    def _mixed(self, runs_data, *, before=0, after=60, line=276, align=None):
        """Paragraph with multiple differently-formatted runs."""
        p = self.doc.add_paragraph()
        self._sp(p, before=before, after=after, line=line)
        if align:
            p.alignment = align
        for rd in runs_data:
            self._run(p, rd["text"],
                      size=rd.get("size", 11),
                      bold=rd.get("bold", False),
                      italic=rd.get("italic", False),
                      color=rd.get("color"))
        return p

    def _section_heading(self, text: str):
        p = self.doc.add_paragraph()
        self._sp(p, before=120, after=20)
        self._run(p, text.upper(), size=13, bold=True, color=self.NAVY)
        self._bottom_border(p)
        return p

    def _job_title_line(self, title: str, period: str):
        p = self.doc.add_paragraph()
        self._sp(p, before=120, after=20)
        self._right_tab(p)
        self._run(p, title,  bold=True)
        self._run(p, "\t")
        self._run(p, period, italic=True)
        return p

    def _company_line(self, company: str, location: str):
        p = self.doc.add_paragraph()
        self._sp(p, before=0, after=40, line=256)
        self._right_tab(p)
        self._run(p, company,  bold=True, color=self.NAVY)
        self._run(p, "\t")
        self._run(p, location, italic=True, color=self.GRAY)
        return p

    def _bullet(self, text: str):
        """Hanging-indent bullet: left=360tw, hanging=260tw, tab stop at 360."""
        p = self.doc.add_paragraph()
        self._sp(p, before=0, after=20, line=256)
        fmt = p.paragraph_format
        fmt.left_indent        = Twips(360)
        fmt.first_line_indent  = Twips(-260)
        # Left tab stop at indent position so wrapped lines align correctly
        pPr  = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab  = OxmlElement("w:tab")
        tab.set(qn("w:val"), "left")
        tab.set(qn("w:pos"), "360")
        tabs.append(tab)
        pPr.append(tabs)
        self._run(p, "\u2022\t")
        self._run(p, " ".join(text.split()))
        return p

    def _project_header(self, text: str):
        p = self.doc.add_paragraph()
        self._sp(p, before=60, after=20, line=256)
        self._run(p, text, bold=True, italic=True, color=self.GRAY)
        return p

    # ── document build ────────────────────────────────────────────────

    def _build(self):
        pers = self.cv.get("personal", {})

        # Header
        self._para(pers.get("name", ""), size=26, bold=True, color=self.NAVY,
                   align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=40)
        if pers.get("title"):
            self._para(pers["title"], size=12, color=self.GRAY,
                       align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
        contact = [x for x in [
            pers.get("location"),
            pers.get("phone"),
            pers.get("email"),
            f"github.com/{pers['github']}"    if pers.get("github")   else None,
            f"linkedin.com/in/{pers['linkedin']}" if pers.get("linkedin") else None,
        ] if x]
        if contact:
            self._para("  \u2022  ".join(contact), size=10, color=self.GRAY,
                       align=WD_ALIGN_PARAGRAPH.CENTER, after=120)

        # Profile
        if self.cv.get("profile"):
            self._section_heading("Profile")
            self._para(self.cv["profile"].strip(), after=80)

        # Technical Skills
        if self.cv.get("skills"):
            self._section_heading("Technical Skills")
            for s in self.cv["skills"]:
                self._mixed([
                    {"text": s["category"] + ": ", "bold": True, "color": "1F3A5F"},
                    {"text": str(s.get("items", ""))}
                ], after=20)

        # Experience
        if self.cv.get("experience"):
            self._section_heading("Experience")
            for job in self.cv["experience"]:
                self._job_title_line(
                    job.get("title", ""),
                    job.get("period", "")
                )
                self._company_line(
                    job.get("company", ""),
                    job.get("location", "")
                )
                for h in job.get("highlights", []):
                    if h.get("title"):
                        self._project_header(h["title"])
                    for pt in h.get("points", []):
                        self._bullet(" ".join(pt.split()))
                for pt in job.get("points", []):
                    self._bullet(" ".join(pt.split()))

        # Education
        if self.cv.get("education"):
            self._section_heading("Education")
            for e in self.cv["education"]:
                suffix = f" ({e['year']})" if e.get("year") else ""
                self._mixed([
                    {"text": e.get("degree", ""), "bold": True},
                    {"text": f" \u2014 {e.get('institution', '')}{suffix}",
                     "color": "595959"}
                ], after=0)

    def get_bytes(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


# ── Routes ────────────────────────────────────────────────────────────

@app.route("/")
def index():
    cv = load_cv()
    return render_template("index.html", cv=cv)


@app.route("/download/txt")
def download_txt():
    cv = load_cv()
    content = cv_to_markdown(cv)
    name = cv.get("personal", {}).get("name", "resume").replace(" ", "_")
    return Response(
        content,
        mimetype="text/plain; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={name}_CV.txt"}
    )


@app.route("/download/pdf")
def download_pdf():
    cv = load_cv()
    pdf = CvPDF(cv)
    buf = io.BytesIO(pdf.output())
    name = cv.get("personal", {}).get("name", "Resume").replace(" ", "_")
    return Response(
        buf.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={name}_CV.pdf"}
    )


@app.route("/download/docx")
def download_docx():
    cv = load_cv()
    doc = CvDOCX(cv)
    name = cv.get("personal", {}).get("name", "Resume").replace(" ", "_")
    return Response(
        doc.get_bytes(),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={name}_CV.docx"}
    )


@app.route("/generate")
def generate():
    return render_template("generate.html")


@app.route("/generate/stream", methods=["POST"])
def generate_stream():
    role_title    = request.form.get("role_title", "").strip()
    company       = request.form.get("company", "").strip()
    job_description = request.form.get("job_description", "").strip()

    def error_sse(msg: str):
        return Response(
            f"data: {json.dumps({'error': msg})}\n\ndata: [DONE]\n\n",
            mimetype="text/event-stream",
            headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"}
        )

    if not role_title or not job_description:
        return error_sse("Role title and job description are required.")

    api_key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        return error_sse("ANTHROPIC_API_KEY is not configured on the server.")

    cv = load_cv()
    resume_md = cv_to_markdown(cv)

    user_message = (
        f"Here is my current CV in Markdown format:\n\n---\n{resume_md}\n---\n\n"
        f"Please tailor this CV for the following position:\n\n"
        f"**Role Title:** {role_title}\n"
        f"**Company:** {company or 'Not specified'}\n\n"
        f"**Job Description / Requirements:**\n{job_description}\n\n"
        "Generate the tailored CV in clean Markdown. "
        "Start directly with the name/header — no preamble."
    )

    def sse_stream():
        try:
            client = anthropic.Anthropic(api_key=api_key)
            with client.messages.stream(
                model="claude-opus-4-6",
                max_tokens=4096,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": user_message}]
            ) as stream:
                for chunk in stream.text_stream:
                    yield f"data: {json.dumps({'token': chunk})}\n\n"
        except anthropic.APIError as exc:
            yield f"data: {json.dumps({'error': str(exc)})}\n\n"
        finally:
            yield "data: [DONE]\n\n"

    return Response(
        stream_with_context(sse_stream()),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"}
    )


if __name__ == "__main__":
    port  = int(os.environ.get("PORT", 5001))
    debug = os.environ.get("FLASK_DEBUG", "false").lower() == "true"
    app.run(host="0.0.0.0", port=port, debug=debug)
